#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
sync_xs.py — 把用友云毛利明细 xlsx 转成 TSV 并写入底表 XS 区（恢复旧 SOP 断点）

背景：
    旧 SOP = 拉数 → update_xs.py 写底表 XS → 公式自动重算 → 看板读底表口径。
    9月管线自动化改造时漏接了写 XS 这一步，导致底表 XS 滞后、
    底表与看板数据不一致（晨哥 2026-09-10 发现并拍板恢复）。

本脚本职责：
    1. 读用友云导出的毛利明细 xlsx（19列）
    2. 补齐 S列销售成本（缺失时 = 金额M − 毛利N，与 8 月 SOP 口径一致）
    3. 生成符合 update_xs.py EXPECT_HEADER 的 TSV
    4. 调用 update_xs.py 驱动本机 Excel 写入底表 XS 并全量重算

用法：
    python sync_xs.py <毛利明细.xlsx> [--update-xs 路径]
"""
import argparse
import os
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

EXPECT_HEADER = [
    "出库单号", "单据类型", "出库日期", "商品分类", "商品sku分类",
    "商品SKU编码", "商品名称", "入库属性", "数量", "单价", "原价",
    "折扣价", "金额", "毛利", "SO激励", "业务员", "库区",
    "销售出库单门店", "销售成本",
]
STORE_NAME = "华为李家村万达授权体验店"


def clean_num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return ""
    try:
        f = float(str(v).replace(",", "").strip())
        return str(int(f)) if f == int(f) else str(f)
    except (TypeError, ValueError):
        return str(v).strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("maoli_xlsx", help="用友云毛利明细xlsx")
    ap.add_argument("--update-xs", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "update_xs.py"))
    args = ap.parse_args()

    wb = load_workbook(args.maoli_xlsx, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        r = list(r[:19]) + [""] * max(0, 19 - len(r))
        # S列销售成本：缺失时 = 金额(12) − 毛利(13)，保留已有的真实值
        if r[18] in (None, ""):
            try:
                cost = float(r[12] or 0) - float(r[13] or 0)
                r[18] = cost
            except (TypeError, ValueError):
                pass
        rows.append(r)

    if not rows:
        sys.exit("❌ 明细无数据行，中止写底表")

    # 汇总校验信息（写入前留证）
    gross = sum(float(r[13] or 0) for r in rows if clean_num(r[13]))
    print(f"→ 明细 {len(rows)} 行，毛利合计 ¥{gross:,.2f}")

    # 生成 TSV
    tsv = tempfile.NamedTemporaryFile(
        mode="w", suffix=".tsv", delete=False, encoding="utf-8",
        dir=os.path.dirname(os.path.abspath(args.maoli_xlsx)),
    )
    with tsv as f:
        f.write("\t".join(EXPECT_HEADER) + "\n")
        for r in rows:
            cells = []
            for i, v in enumerate(r):
                if i == 2:  # 出库日期去时间部分
                    cells.append(str(v).replace(" 00:00:00", "").strip())
                elif i in (8, 9, 10, 11, 12, 13, 14, 18):
                    cells.append(clean_num(v))
                else:
                    cells.append(str(v if v is not None else "").strip())
            f.write("\t".join(cells) + "\n")
    print(f"→ TSV 就绪: {tsv.name}")

    # 调 update_xs.py 写底表（其内部自带备份/表头校验/Excel驱动）
    res = subprocess.run(
        [sys.executable, args.update_xs, tsv.name],
        capture_output=True, text=True,
    )
    os.unlink(tsv.name)
    print(res.stdout.strip())
    if res.returncode != 0:
        print(res.stderr.strip()[:800], file=sys.stderr)
        sys.exit("❌ update_xs 写底表失败")


if __name__ == "__main__":
    main()
