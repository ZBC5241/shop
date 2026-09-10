#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
write_lehuishou_xml.py — XML级更新底表「乐机收」T/U列（update_lehuishou.py 的XML版）

背景：
    update_lehuishou.py 用 openpyxl 整本保存，会丢条件格式样式（dxf）；
    本脚本直接原位替换「李家村销售」sheet XML 中 T14:U18 的8个单元格，
    其余内容（公式/条件格式/样式）零损伤。

单元格映射（与 calc_data.py 完全同源）：
    行 = 月任务行号 + 10；张博晨固定 18
    T列=单量，U列=公司净利；V列增值率（备用，不写）
    T19/U19 是 SUM 公式，不动。

用法：
    python write_lehuishou_xml.py --邵乐乐 4 217 --杨丽华 5 474 --李泽 10 509 --张博晨 1 223
    （数值为 0 也要显式传 0，会覆盖为 0）
"""
import argparse
import datetime
import os
import re
import shutil
import sys
import tempfile
import zipfile

DEFAULT_XLSX = "/Users/mac/Desktop/李家村销售/李家村9月任务进度.xlsx"
TARGET_SHEET_FILE = "xl/worksheets/sheet3.xml"  # 李家村销售 = 第3个sheet

# 与 calc_data.py 同源：月任务行号+10，张博晨固定18
PERSON_ROW = {"邵乐乐": 14, "杨丽华": 15, "李泽": 16, "张博晨": 18}


def cell_xml(ref, v, t="n", s=None):
    """生成 <c> 元素：t='n' 数值 / t='str' 字符串。"""
    attrs = f' s="{s}"' if s is not None else ""
    if v is None:
        return f'<c r="{ref}"{attrs} t="{t}" />'
    if t == "n":
        return f'<c r="{ref}"{attrs} t="n"><v>{v}</v></c>'
    return f'<c r="{ref}"{attrs} t="str"><is><t>{v}</t></is></c>'


def main():
    ap = argparse.ArgumentParser()
    for n in PERSON_ROW:
        ap.add_argument(f"--{n}", nargs=2, type=float, metavar=("单量", "净利"))
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    args = ap.parse_args()

    updates = {}
    for n in PERSON_ROW:
        v = getattr(args, n)
        if v is not None:
            updates[PERSON_ROW[n]] = (n, v[0], v[1])
    if not updates:
        sys.exit("❌ 未提供任何人员数据")

    target = args.xlsx
    if not os.path.exists(target):
        sys.exit(f"❌ 找不到底表: {target}")

    # 读取当前 T/U 旧值（留证）
    from openpyxl import load_workbook
    wb_old = load_workbook(target, read_only=True)
    ws_old = wb_old["李家村销售"]
    for r in sorted(updates):
        name, o, a = updates[r]
        old_t = ws_old.cell(r, 20).value
        old_u = ws_old.cell(r, 21).value
        flag = "→" if (old_t != o or old_u != a) else "=（无变化）"
        print(f"  {name}(行{r}): T {old_t}→{o:.0f}  U {old_u}→{a:g} {flag}")

    need_write = any(ws_old.cell(r, 20).value != updates[r][1] or
                     ws_old.cell(r, 21).value != updates[r][2]
                     for r in updates)
    if not need_write:
        print("✓ 所有值已是最新，无需写入")
        return

    # 备份（保留10份）
    bak_dir = os.path.join(os.path.dirname(target), "_备份")
    os.makedirs(bak_dir, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    base = os.path.splitext(os.path.basename(target))[0]
    bak = os.path.join(bak_dir, f"{base}_备份{stamp}.xlsx")
    shutil.copy2(target, bak)
    print(f"→ 已备份: {os.path.basename(bak)}")
    for old in sorted(f for f in os.listdir(bak_dir) if f.startswith(base + "_备份"))[:-10]:
        os.remove(os.path.join(bak_dir, old))

    # XML 原位替换：只在「李家村销售」sheet XML 内替换 T/U 单元格
    with zipfile.ZipFile(target, "r") as zin:
        xml = zin.read(TARGET_SHEET_FILE).decode("utf-8")

    for r, (name, o, a) in sorted(updates.items()):
        for col, val in (("T", o), ("U", a)):
            ref = f"{col}{r}"
            # 匹配该单元格现有XML（自闭合或带值），替换为带值版本
            # 属性顺序为 s=... t=...（中间留空格），自闭合 `<c .../>` 或带值 `<c ...>...</c>`
            pat = re.compile(
                r'<c r="%s"(?: s="(\d+)")?(?: t="[^"]+")?\s*(?:/>|>(?:(?!</c>|<c ).)*</c>)' % ref,
                re.S,
            )
            m = pat.search(xml)
            if not m:
                sys.exit(f"❌ 未找到单元格 {ref}，中止")
            style = m.group(1)
            new_cell = cell_xml(ref, int(val) if float(val) == int(val) else val,
                                t="n", s=style)
            xml = pat.sub(new_cell, xml, count=1)

    tmp_fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(tmp_fd)
    with zipfile.ZipFile(target, "r") as zin, \
         zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == TARGET_SHEET_FILE:
                data = xml.encode("utf-8")
            elif item.filename == "xl/workbook.xml":
                x = data.decode("utf-8")
                if "fullCalcOnLoad" not in x:
                    if "<calcPr" in x:
                        x = re.sub(r"<calcPr", '<calcPr fullCalcOnLoad="1"', x, count=1)
                    else:
                        x = x.replace("</workbook>", '<calcPr calcId="0" fullCalcOnLoad="1"/></workbook>')
                data = x.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp_path, target)
    print("→ XML 原位替换完成")

    # 校验
    wb2 = load_workbook(target)
    ws2 = wb2["李家村销售"]
    ok = True
    for r, (name, o, a) in updates.items():
        gt, gu = ws2.cell(r, 20).value, ws2.cell(r, 21).value
        if gt != o or gu != a:
            ok = False
            print(f"  ❌ {name}: T={gt} U={gu} 期望 {o}/{a}")
    assert ok, "写入校验失败"
    # 公式完好性
    assert str(ws2["T19"].value).startswith("=SUM"), "T19公式丢失"
    assert str(ws2["U19"].value).startswith("=SUM"), "U19公式丢失"
    aa = ws2["AA14"].value
    aa_text = getattr(aa, "text", aa)
    assert "SUMIFS" in str(aa_text), "AA14公式丢失"
    t_sum = sum(ws2.cell(r, 20).value or 0 for r in range(14, 19) if isinstance(ws2.cell(r, 20).value, (int, float)))
    u_sum = sum(ws2.cell(r, 21).value or 0 for r in range(14, 19) if isinstance(ws2.cell(r, 21).value, (int, float)))
    print(f"✓ 校验通过: 单量合计 {t_sum} | 净利合计 ¥{u_sum:g} | 公式完好")


if __name__ == "__main__":
    main()
